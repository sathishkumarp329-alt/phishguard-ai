"""Word Document (.docx) URL Extraction Module."""

import io
import re
from typing import List, Set
import docx

URL_PATTERN = re.compile(
    r"https?://(?:[a-zA-Z0-9-._~:/?#\[\]@!$&'()*+,;=]|%[0-9a-fA-F]{2})+"
)


def _clean_url(url: str) -> str:
    return url.rstrip(".,;:)'\"!>?]")


def extract_urls_from_docx(content: bytes) -> List[str]:
    """
    Extract URLs from Word documents (.docx), scanning paragraph text,
    table cells, and internal XML relationship hyperlinks.
    """
    extracted_urls: List[str] = []
    seen: Set[str] = set()

    def _add_url(u: str) -> None:
        clean = _clean_url(u.strip())
        if clean and clean.startswith(("http://", "https://")) and clean not in seen:
            seen.add(clean)
            extracted_urls.append(clean)

    try:
        doc = docx.Document(io.BytesIO(content))

        # 1. Hyperlinks embedded in document relationships
        try:
            for rel in doc.part.rels.values():
                if "hyperlink" in rel.reltype.lower():
                    _add_url(rel.target_ref)
        except Exception:
            pass

        # 2. Text in paragraphs
        for p in doc.paragraphs:
            if p.text:
                for match in URL_PATTERN.findall(p.text):
                    _add_url(match)

        # 3. Text in tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        for match in URL_PATTERN.findall(cell.text):
                            _add_url(match)

    except Exception:
        # Corrupted or unreadable DOCX
        pass

    return extracted_urls
